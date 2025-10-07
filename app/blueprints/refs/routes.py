import re
from flask import session, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import io
from flask import render_template, request, redirect, url_for, flash, jsonify, Response
from flask_login import login_required
from sqlalchemy.exc import IntegrityError
from . import bp
from .forms import RefCodeForm, QuickTypeForm
from ...extensions import db
from ...models import RefCode, CodeType


@bp.route("/codes", methods=["GET"])
@login_required
def codes_list():
    # --- параметры ---
    q = request.args.get("q", "", type=str).strip()
    status = request.args.get("status", "", type=str).strip()  # "", "active", "inactive"
    type_id = request.args.get("type_id", 0, type=int)
    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 20, type=int)

    # --- базовый запрос ---
    query = RefCode.query.join(CodeType)

    if q:
        like = f"%{q}%"
        query = query.filter(
            db.or_(
                RefCode.code.ilike(like),
                RefCode.name.ilike(like),
                CodeType.name.ilike(like),
            )
        )

    if status in ("active", "inactive"):
        query = query.filter(RefCode.status == status)

    if type_id:
        query = query.filter(RefCode.type_id == type_id)

    query = query.order_by(RefCode.created_at.desc())

    pagination = query.paginate(page=page, per_page=per_page, error_out=False)
    items = pagination.items

    # список типов для фильтра
    types = CodeType.query.order_by(CodeType.name.asc()).all()

    return render_template(
        "refs/codes_list.html",
        items=items,
        pagination=pagination,
        q=q,
        status=status,
        type_id=type_id,
        per_page=per_page,
        types=types,
    )


@bp.route("/codes/new", methods=["GET", "POST"])
@login_required
def codes_new():
    form = RefCodeForm()
    form.set_type_choices()
    if request.method == "POST" and form.validate_on_submit():
        item = RefCode(
            code=form.code.data,
            name=form.name.data,
            type_id=form.type_id.data,
            status=form.status.data,
        )
        db.session.add(item)
        try:
            db.session.commit()
            flash("Код успешно добавлен", "success")
            return redirect(url_for("refs.codes_list"))
        except IntegrityError:
            db.session.rollback()
            flash("Такой код уже существует (уникальность по полю Код).", "warning")
    return render_template("refs/codes_new.html", form=form)


@bp.route("/codes/<int:item_id>/edit", methods=["GET", "POST"])
@login_required
def codes_edit(item_id: int):
    item = RefCode.query.get_or_404(item_id)
    form = RefCodeForm(obj=item)
    form.set_type_choices()
    if request.method == "POST" and form.validate_on_submit():
        item.code = form.code.data
        item.name = form.name.data
        item.type_id = form.type_id.data
        item.status = form.status.data
        try:
            db.session.commit()
            flash("Изменения сохранены", "success")
            return redirect(url_for("refs.codes_list"))
        except IntegrityError:
            db.session.rollback()
            flash("Нарушена уникальность кода (дубликат).", "warning")
    return render_template("refs/codes_edit.html", form=form, item=item)


@bp.route("/codes/<int:item_id>/toggle", methods=["POST"])
@login_required
def codes_toggle(item_id: int):
    item = RefCode.query.get_or_404(item_id)
    item.status = "inactive" if item.status == "active" else "active"
    db.session.commit()
    flash("Статус переключён", "info")
    return redirect(url_for("refs.codes_list", **request.args))


@bp.route("/codes/<int:item_id>/delete", methods=["POST"])
@login_required
def codes_delete(item_id: int):
    item = RefCode.query.get_or_404(item_id)
    try:
        db.session.delete(item)
        db.session.commit()
        flash("Код удалён", "success")
    except IntegrityError:
        db.session.rollback()
        flash("Невозможно удалить: есть связанные записи.", "warning")
    # сохраняем текущие фильтры/страницу при возврате
    return redirect(url_for("refs.codes_list", **request.args))


# Быстрое добавление типа (AJAX)
@bp.route("/code-types/quick-add", methods=["POST"])
@login_required
def code_types_quick_add():
    form = QuickTypeForm()
    if form.validate_on_submit():
        name = form.name.data.strip()
        if not name:
            return jsonify({"ok": False, "error": "Пустое имя"}), 400

        exists = CodeType.query.filter(db.func.lower(CodeType.name) == name.lower()).first()
        if exists:
            return jsonify({"ok": True, "id": exists.id, "name": exists.name, "exists": True}), 200

        t = CodeType(name=name)
        db.session.add(t)
        db.session.commit()
        return jsonify({"ok": True, "id": t.id, "name": t.name}), 201

    return jsonify({"ok": False, "errors": form.errors}), 400

@bp.route("/codes/export.xlsx", methods=["GET"])
@login_required
def codes_export_xlsx():
    # Параметры фильтра  как в списке
    q = request.args.get("q", "", type=str).strip()
    status = request.args.get("status", "", type=str).strip()
    type_id = request.args.get("type_id", 0, type=int)

    query = RefCode.query.join(CodeType)
    if q:
        like = f"%{q}%"
        query = query.filter(
            db.or_(
                RefCode.code.ilike(like),
                RefCode.name.ilike(like),
                CodeType.name.ilike(like),
            )
        )
    if status in ("active", "inactive"):
        query = query.filter(RefCode.status == status)
    if type_id:
        query = query.filter(RefCode.type_id == type_id)

    query = query.order_by(RefCode.created_at.desc())
    rows = query.all()

    # Создаём книгу
    wb = Workbook()
    ws = wb.active
    ws.title = "Коды"

    headers = ["Код", "Наименование", "Тип", "Статус", "Создан", "Обновлён"]
    ws.append(headers)

    # Стили заголовка
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E9ECEF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="D0D0D0")
    header_border = Border(top=thin, left=thin, right=thin, bottom=thin)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border

    # Данные
    for r in rows:
        ws.append([
            r.code,
            r.name,
            r.type.name if r.type else "",
            "Активен" if r.status == "active" else "Не активен",
            r.created_at,
            r.updated_at,
        ])

    # Форматы и автоширина
    date_fmt = "yyyy-mm-dd hh:mm"
    max_lens = [len(h) for h in headers]

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=6):
        for idx, cell in enumerate(row, start=1):
            # даты
            if idx in (5, 6) and cell.value:
                cell.number_format = date_fmt
            # под автоширину
            text = ""
            val = cell.value
            if val is None:
                text = ""
            elif isinstance(val, str):
                text = val
            else:
                text = str(val)
            if len(text) > max_lens[idx - 1]:
                max_lens[idx - 1] = len(text)

    # Устанавливаем ширины столбцов (мин/макс, чтобы красиво)
    for i, m in enumerate(max_lens, start=1):
        width = min(max(m + 2, 10), 60)
        ws.column_dimensions[get_column_letter(i)].width = width

    # Заморозка заголовка
    ws.freeze_panes = "A2"

    # Выгрузка в байтовый буфер
    buff = io.BytesIO()
    wb.save(buff)
    data = buff.getvalue()
    buff.close()

    filename = "ref_codes_export.xlsx"
    headers_resp = {
        "Content-Disposition": f'attachment; filename="{filename}"'
    }
    return Response(
        data,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_resp,
    )

@bp.route("/codes/import", methods=["GET", "POST"])
@login_required
def codes_import():
    """
    Загрузка xlsx -> превью (без записи в БД).
    Ожидаемые столбцы (имена в первой строке):
      - Код
      - Наименование
      - Тип
      - Статус
    """
    create_types = bool(request.args.get("create_types") == "1" or request.form.get("create_types") == "1")

    if request.method == "GET":
        session.pop("codes_import_preview", None)
        return render_template("refs/codes_import.html", preview=None, create_types=create_types)

    f = request.files.get("file")
    if not f or f.filename.strip() == "":
        flash("Файл не выбран", "warning")
        return render_template("refs/codes_import.html", preview=None, create_types=create_types)

    try:
        from openpyxl import load_workbook
    except Exception:
        flash("Не установлена зависимость openpyxl", "danger")
        return redirect(url_for("refs.codes_list"))

    try:
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.active
    except Exception as e:
        flash(f"Не удалось прочитать Excel: {e}", "danger")
        return render_template("refs/codes_import.html", preview=None, create_types=create_types)

    # Заголовки
    headers = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        headers.append((cell.value or "").strip() if isinstance(cell.value, str) else str(cell.value or "").strip())

    def find_idx(name):
        name_low = name.lower()
        for i, h in enumerate(headers):
            if (h or "").strip().lower() == name_low:
                return i
        return None

    idx_code = find_idx("код")
    idx_name = find_idx("наименование")
    idx_type = find_idx("тип")
    idx_status = find_idx("статус")

    missing = [label for label, idx in [("Код", idx_code), ("Наименование", idx_name), ("Тип", idx_type), ("Статус", idx_status)] if idx is None]
    if missing:
        flash("В первой строке Excel должны быть столбцы: Код, Наименование, Тип, Статус", "danger")
        return render_template("refs/codes_import.html", preview=None, create_types=create_types)

    # Справочники
    all_types = {t.name.lower(): t for t in CodeType.query.all()}
    existing_codes = {c.code for c in RefCode.query.all()}

    re_code = re.compile(r"^\d{5}$")

    def norm_status(s):
        if not s: return None
        s = str(s).strip().lower()
        mapping = {
            "активен": "active", "active": "active", "да": "active", "1": "active", "true": "active",
            "не активен": "inactive", "неактивен": "inactive", "inactive": "inactive", "нет": "inactive", "0": "inactive", "false": "inactive",
        }
        return mapping.get(s)

    rows = []
    errors = 0
    duplicates = 0
    duplicates_in_file = 0
    new_types = set()
    seen_in_file = set()

    for r in ws.iter_rows(min_row=2, values_only=True):
        code_val = (r[idx_code] or "").strip() if isinstance(r[idx_code], str) else str(r[idx_code] or "").strip()
        name_val = (r[idx_name] or "").strip() if isinstance(r[idx_name], str) else str(r[idx_name] or "").strip()
        type_val = (r[idx_type] or "").strip() if isinstance(r[idx_type], str) else str(r[idx_type] or "").strip()
        status_val = norm_status(r[idx_status])

        row_status = []
        ok = True

        if not re_code.match(code_val):
            ok = False
            row_status.append("Код не 5-значный")

        if not name_val:
            ok = False
            row_status.append("Пустое наименование")

        if not type_val:
            ok = False
            row_status.append("Пустой тип")
        else:
            if type_val.lower() not in all_types:
                new_types.add(type_val)
                if not create_types:
                    row_status.append("Новый тип (будет ошибка без разрешения)")

        if not status_val:
            ok = False
            row_status.append("Статус не распознан")

        # дубликаты в БД
        if code_val in existing_codes:
            ok = False
            row_status.append("Дубликат в БД")
            duplicates += 1

        # дубликаты внутри файла
        if code_val in seen_in_file:
            ok = False
            row_status.append("Дубликат в файле")
            duplicates_in_file += 1
        seen_in_file.add(code_val)

        rows.append({
            "code": code_val,
            "name": name_val,
            "type": type_val,
            "status": status_val,
            "ok": ok,
            "problems": row_status
        })
        if not ok:
            errors += 1

    preview = {
        "total": len(rows),
        "errors": errors,
        "duplicates": duplicates,
        "duplicates_in_file": duplicates_in_file,
        "new_types": sorted(new_types),
        "rows": rows[:200],
        "create_types": create_types,
    }

    session["codes_import_preview"] = {
        "create_types": create_types,
        "rows": rows,  # полностью (без усечения)
        "new_types": sorted(new_types),
    }

    return render_template("refs/codes_import.html", preview=preview, create_types=create_types)


@bp.route("/codes/import/commit", methods=["POST"])
@login_required
def codes_import_commit():
    data = session.get("codes_import_preview")
    if not data:
        flash("Нет подготовленных данных импорта. Загрузите файл ещё раз.", "warning")
        return redirect(url_for("refs.codes_import"))

    create_types = bool(request.form.get("create_types") == "1" or data.get("create_types"))
    rows = data.get("rows", [])
    if not rows:
        flash("Пустые данные импорта.", "warning")
        return redirect(url_for("refs.codes_import"))

    all_types = {t.name.lower(): t for t in CodeType.query.all()}
    existing_codes = {c.code for c in RefCode.query.all()}

    created_types_count = 0
    created_codes_count = 0
    skipped_count = 0

    results = []  # для отчёта: [code, name, type, status_text, result, reason]

    # Создадим недостающие типы заранее (если разрешено)
    if create_types:
        for tname in data.get("new_types", []):
            key = tname.lower()
            if key not in all_types:
                t = CodeType(name=tname)
                db.session.add(t)
                db.session.flush()
                all_types[key] = t
                created_types_count += 1

    for r in rows:
        code_val = r["code"]; name_val = r["name"]; type_val = r["type"]; status_val = r["status"]
        status_text = "Активен" if status_val == "active" else ("Не активен" if status_val == "inactive" else str(status_val))

        if not r.get("ok"):
            skipped_count += 1
            results.append([code_val, name_val, type_val, status_text, "SKIPPED", "; ".join(r.get("problems", []))])
            continue

        if code_val in existing_codes:
            skipped_count += 1
            results.append([code_val, name_val, type_val, status_text, "SKIPPED", "Дубликат в БД"])
            continue

        key = type_val.lower()
        t = all_types.get(key)
        if not t:
            if not create_types:
                skipped_count += 1
                results.append([code_val, name_val, type_val, status_text, "SKIPPED", "Новый тип запрещён"])
                continue
            t = CodeType(name=type_val)
            db.session.add(t)
            db.session.flush()
            all_types[key] = t
            created_types_count += 1

        item = RefCode(code=code_val, name=name_val, type_id=t.id, status=status_val)
        db.session.add(item)
        existing_codes.add(code_val)
        created_codes_count += 1
        results.append([code_val, name_val, type_val, status_text, "OK", ""])

    db.session.commit()
    session.pop("codes_import_preview", None)

    # --- Формируем Excel-лог ---
    wb = Workbook()
    ws1 = wb.active; ws1.title = "Summary"
    ws2 = wb.create_sheet("Rows")

    # Summary
    ws1.append(["Metric", "Value"])
    ws1.append(["Created codes", created_codes_count])
    ws1.append(["Created types", created_types_count])
    ws1.append(["Skipped", skipped_count])

    # Rows header
    hdr = ["Код", "Наименование", "Тип", "Статус", "Результат", "Причина"]
    ws2.append(hdr)
    for row in results:
        ws2.append(row)

    # Небольшое форматирование
    from openpyxl.styles import Font, PatternFill, Alignment
    head_fill = PatternFill("solid", fgColor="E9ECEF")
    head_font = Font(bold=True)
    for c in ws2[1]:
        c.font = head_font; c.fill = head_fill; c.alignment = Alignment(horizontal="center")

    # Автоширина
    from openpyxl.utils import get_column_letter
    for sheet in (ws1, ws2):
        max_cols = sheet.max_column
        for col in range(1, max_cols + 1):
            max_len = 10
            for cell in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=col, max_col=col):
                v = cell[0].value
                s = v if isinstance(v, str) else ("" if v is None else str(v))
                if len(s) > max_len: max_len = len(s)
            sheet.column_dimensions[get_column_letter(col)].width = min(max_len + 2, 60)

    buff = io.BytesIO()
    wb.save(buff); buff.seek(0)

    return send_file(buff,
                     as_attachment=True,
                     download_name="ref_codes_import_result.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/codes/import/template.xlsx", methods=["GET"])
@login_required
def codes_import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Шаблон"
    headers = ["Код", "Наименование", "Тип", "Статус"]
    ws.append(headers)
    ws.append(["12345", "Пример позиции", "Поставщик", "Активен"])
    ws.append(["67890", "Ещё пример", "Клиент", "Не активен"])

    # стили заголовка
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E9ECEF")
    header_align = Alignment(horizontal="center", vertical="center")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align

    # автоширина
    from openpyxl.utils import get_column_letter
    max_lens = [len(h) for h in headers]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=4):
        for idx, cell in enumerate(row, start=1):
            val = cell.value
            s = val if isinstance(val, str) else ("" if val is None else str(val))
            if len(s) > max_lens[idx - 1]:
                max_lens[idx - 1] = len(s)
    for i, m in enumerate(max_lens, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(m + 2, 10), 50)

    buff = io.BytesIO()
    wb.save(buff)
    buff.seek(0)
    return send_file(buff, as_attachment=True, download_name="ref_codes_template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@bp.route("/codes/batch", methods=["POST"])
@login_required
def codes_batch():
    ids = request.form.getlist("ids")
    action = request.form.get("action", "").strip()
    if not ids:
        flash("Ничего не выбрано", "warning")
        return redirect(url_for("refs.codes_list", **request.args))

    q = RefCode.query.filter(RefCode.id.in_([int(x) for x in ids]))
    items = q.all()

    affected = 0
    if action == "activate":
        for it in items:
            if it.status != "active":
                it.status = "active"; affected += 1
        db.session.commit()
        flash(f"Активировано: {affected}", "success")
    elif action == "deactivate":
        for it in items:
            if it.status != "inactive":
                it.status = "inactive"; affected += 1
        db.session.commit()
        flash(f"Деактивировано: {affected}", "success")
    elif action == "delete":
        deleted = 0; failed = 0
        for it in items:
            try:
                db.session.delete(it); db.session.flush()
                deleted += 1
            except IntegrityError:
                db.session.rollback()
                failed += 1
        db.session.commit()
        msg = f"Удалено: {deleted}"
        if failed:
            msg += f", не удалось: {failed}"
        flash(msg, "success" if failed == 0 else "warning")
    else:
        flash("Неизвестное действие", "warning")

    return redirect(url_for("refs.codes_list", **request.args))

