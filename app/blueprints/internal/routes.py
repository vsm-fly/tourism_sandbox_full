# app/blueprints/internal/routes.py
from __future__ import annotations

from io import BytesIO
from datetime import date

from flask import render_template, redirect, url_for, flash, request, current_app, send_file, abort
from flask_login import login_required, current_user
from sqlalchemy import or_

from app.extensions import db
from app.models import InternalTour, RefCode
from app.models.party import Client  # импорт напрямую из party
from app.forms.internal_tour import InternalTourForm
from . import bp


# ---------- утилиты ----------

def _require_admin():
    if not getattr(current_user, "is_admin", False):
        r = getattr(current_user, "role", None)
        role_val = ((getattr(r, "name", None) or getattr(r, "code", None) or "") if r else "").strip().lower()
        if role_val != "admin":
            abort(403)


def _safe_set(obj, primary: str, value, *aliases: str) -> bool:
    """Безопасно проставляет значение в первый доступный атрибут."""
    for name in (primary, *aliases):
        if not hasattr(obj, name):
            continue
        try:
            setattr(obj, name, value)
            return True
        except Exception:
            continue
    return False


def _build_kind_labels() -> dict:
    form = InternalTourForm()
    choices = getattr(form, "kind", None)
    choices = getattr(choices, "choices", []) or []
    mapping: dict = {}
    for k, v in choices:
        mapping[k] = v
        mapping[str(k).strip().lower()] = v
    fallback = {"internal": "Тур", "transfer": "Трансфер", "hotel": "Отель", "excursion": "Экскурсия"}
    for k, v in fallback.items():
        mapping.setdefault(k, v)
        mapping.setdefault(str(k).strip().lower(), v)
    return mapping


def _kinds_for_filters():
    form = InternalTourForm()
    choices = getattr(form, "kind", None)
    choices = getattr(choices, "choices", []) or []
    if choices:
        return choices
    return [("internal", "Тур"), ("transfer", "Трансфер"), ("hotel", "Отель"), ("excursion", "Экскурсия")]


def _parse_iso(d: str | None) -> date | None:
    if not d:
        return None
    try:
        return date.fromisoformat(d)
    except Exception:
        return None


def _apply_filters(base_query):
    """
    Поиск только по реальным колонкам InternalTour.
    Прямой фильтр по supplier_code/client_code не делаем (это property).
    """
    q = (request.args.get("q") or "").strip()
    kind = (request.args.get("kind") or "").strip()
    status = (request.args.get("status") or "").strip()

    date_from = _parse_iso(request.args.get("date_from"))
    date_to   = _parse_iso(request.args.get("date_to"))

    query = base_query

    if q:
        like = f"%{q}%"
        conditions = []
        for attr_name in ("fio", "direction", "order_type", "kind"):
            attr = getattr(InternalTour, attr_name, None)
            if attr is not None and hasattr(attr, "ilike"):
                conditions.append(attr.ilike(like))
        if conditions:
            query = query.filter(or_(*conditions))

    if kind and hasattr(InternalTour, "kind"):
        query = query.filter(InternalTour.kind == kind)
    if status and hasattr(InternalTour, "status"):
        query = query.filter(InternalTour.status == status)
    if date_from:
        query = query.filter(InternalTour.order_date >= date_from)
    if date_to:
        query = query.filter(InternalTour.order_date <= date_to)

    return query


def _ensure_client_by_code(code_raw: str | None) -> int | None:
    """
    Возвращает id клиента по коду. Если нет — создаёт на основании RefCode (или самого кода как имени).
    Коммит не делает — только flush.
    """
    code = (code_raw or "").strip()
    if not code:
        return None
    code = code.upper()  # NEW: нормализуем
    existing = Client.query.filter_by(code=code).first()
    if existing:
        return existing.id
    rc = RefCode.query.filter_by(code=code).first()
    name = (rc.name if rc and rc.name else code).strip() or code
    c = Client(code=code, name=name)
    db.session.add(c)
    db.session.flush()  # получить c.id без коммита
    return c.id


# ---------- страницы ----------

@bp.route("/tours")
@login_required
def tours_list():
    per_page = int(request.args.get("per_page", 20))
    page = int(request.args.get("page", 1))

    base = InternalTour.query
    # Эмуляция NULLS LAST: сначала NOT NULL (False -> 0), затем NULL (True -> 1)
    query = _apply_filters(base).order_by(
        (InternalTour.order_date.is_(None)).asc(),
        InternalTour.order_date.desc(),
        InternalTour.id.desc(),
    )
    pagination = query.paginate(page=page, per_page=per_page, error_out=False)

    # Собираем коды для отображения имён через RefCode
    sup_codes, cli_codes = set(), set()
    for it in pagination.items:
        sc_rel = getattr(it, "supplier", None)
        sc = getattr(sc_rel, "code", None) if sc_rel else getattr(it, "supplier_code", None)
        cc_rel = getattr(it, "client", None)
        cc = getattr(cc_rel, "code", None) if cc_rel else getattr(it, "client_code", None)
        if sc: sup_codes.add(str(sc))
        if cc: cli_codes.add(str(cc))

    supplier_name_by_code, client_name_by_code = {}, {}
    if sup_codes:
        rows = RefCode.query.filter(RefCode.code.in_(sup_codes)).all()
        supplier_name_by_code = {str(r.code): (r.name or str(r.code)) for r in rows}
    if cli_codes:
        rows = RefCode.query.filter(RefCode.code.in_(cli_codes)).all()
        client_name_by_code = {str(r.code): (r.name or str(r.code)) for r in rows}

    kind_labels = _build_kind_labels()

    # Подготовка удобных отображаемых полей
    for it in pagination.items:
        raw_code = (getattr(it, "kind", None) or getattr(it, "order_type", "") or "")
        it.kind_label = kind_labels.get(str(raw_code).strip().lower(), raw_code or "-")

        sc_rel = getattr(it, "supplier", None)
        sc = getattr(sc_rel, "code", None) if sc_rel else getattr(it, "supplier_code", None)
        sc = str(sc or "")
        it.supplier_code_display = sc or "-"
        it.supplier_name = supplier_name_by_code.get(sc) or (sc or "-")

        cc_rel = getattr(it, "client", None)
        cc = getattr(cc_rel, "code", None) if cc_rel else getattr(it, "client_code", None)
        cc = str(cc or "")
        it.client_code_display = cc or "-"
        it.client_name = client_name_by_code.get(cc) or (cc or "-")

    kinds = _kinds_for_filters()

    is_admin = bool(getattr(current_user, "is_admin", False))
    if not is_admin:
        r = getattr(current_user, "role", None)
        role_val = ((getattr(r, "name", None) or getattr(r, "code", None) or "") if r else "").strip().lower()
        is_admin = (role_val == "admin")

    return render_template(
        "internal/tours_list.html",
        tours=pagination,
        kinds=kinds,
        is_admin=is_admin,
    )


@bp.route("/tours/export.xlsx")
@login_required
def tours_export():
    base = InternalTour.query
    query = _apply_filters(base).order_by(
        (InternalTour.order_date.is_(None)).asc(),
        InternalTour.order_date.desc(),
        InternalTour.id.desc(),
    )
    items = query.all()

    sup_codes, cli_codes = set(), set()
    for it in items:
        sc_rel = getattr(it, "supplier", None)
        sc = getattr(sc_rel, "code", None) if sc_rel else getattr(it, "supplier_code", None)
        cc_rel = getattr(it, "client", None)
        cc = getattr(cc_rel, "code", None) if cc_rel else getattr(it, "client_code", None)
        if sc: sup_codes.add(str(sc))
        if cc: cli_codes.add(str(cc))

    supplier_name_by_code, client_name_by_code = {}, {}
    if sup_codes:
        rows = RefCode.query.filter(RefCode.code.in_(sup_codes)).all()
        supplier_name_by_code = {str(r.code): (r.name or str(r.code)) for r in rows}
    if cli_codes:
        rows = RefCode.query.filter(RefCode.code.in_(cli_codes)).all()
        client_name_by_code = {str(r.code): (r.name or str(r.code)) for r in rows}

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Внутренний туризм"

    header = [
        "#", "Дата заказа", "Вид", "ФИО",
        "Код пост.", "Поставщик",
        "Код клиента", "Клиент",
        "Направление", "Дата начала", "Дата окончания",
        "Примечание", "Себестоимость", "Продажа", "Нетто", "Маржа (%)",
    ]
    ws.append(header)

    kind_labels = _build_kind_labels()

    for idx, t in enumerate(items, start=1):
        raw_code = (getattr(t, "kind", None) or getattr(t, "order_type", "") or "")
        kind_label = kind_labels.get(str(raw_code).strip().lower(), raw_code or "")

        sc_rel = getattr(t, "supplier", None)
        sc = getattr(sc_rel, "code", None) if sc_rel else getattr(t, "supplier_code", None)
        sc = str(sc or "")
        supplier_name = supplier_name_by_code.get(sc) or ""

        cc_rel = getattr(t, "client", None)
        cc = getattr(cc_rel, "code", None) if cc_rel else getattr(t, "client_code", None)
        cc = str(cc or "")
        client_name = client_name_by_code.get(cc) or ""

        cost = float(getattr(t, "cost", 0) or 0)
        price = float((getattr(t, "price", None) is not None and t.price)
                      or getattr(t, "sale_price", 0)
                      or getattr(t, "sale", 0)
                      or 0)
        net = float((getattr(t, "net", None) is not None and t.net) or (price - cost))
        mrg = round((net / price) * 100, 1) if price else 0.0

        ws.append([
            idx,
            getattr(t, "order_date", None),
            kind_label,
            getattr(t, "fio", "") or "",
            sc,
            supplier_name,
            cc,
            client_name,
            getattr(t, "direction", "") or "",
            getattr(t, "date_start", getattr(t, "start_date", None)),
            getattr(t, "date_end",   getattr(t, "end_date", None)),
            getattr(t, "note", "") or "",
            cost, price, net, mrg,
        ])

    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for c in col:
            val = c.value
            l = len(str(val)) if val is not None else 0
            if l > max_len:
                max_len = l
        ws.column_dimensions[letter].width = min(max_len + 2, 40)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(
        bio,
        as_attachment=True,
        download_name="internal_tours.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/tours/new", methods=["GET", "POST"], endpoint="new_tour")
@login_required
def tour_create():
    _require_admin()
    form = InternalTourForm()
    if form.validate_on_submit():
        obj = InternalTour()

        # даты (пробуем алиасы)
        _safe_set(obj, "order_date", form.order_date.data, "date_order")
        _safe_set(obj, "start_date", form.start_date.data, "date_start")
        _safe_set(obj, "end_date",   form.end_date.data,   "date_end")

        # простые поля
        _safe_set(obj, "kind", form.kind.data)
        _safe_set(obj, "fio",  form.fio.data)
        _safe_set(obj, "direction", form.direction.data)

        # коды -> реальные FK в clients.id (NOT NULL)
        supplier_code = (form.supplier_code.data or "").strip()
        client_code   = (form.client_code.data or "").strip()
        supplier_id = _ensure_client_by_code(supplier_code)
        client_id   = _ensure_client_by_code(client_code)
        if not supplier_id or not client_id:
            flash("Укажите корректные коды поставщика и клиента.", "danger")
            return render_template("internal/tour_form.html", form=form)

        obj.supplier_id = supplier_id
        obj.client_id   = client_id

        # деньги/примечание
        _safe_set(obj, "cost",  form.cost.data or 0)
        _safe_set(obj, "price", form.price.data or 0)
        _safe_set(obj, "note",  form.note.data)

        # пересчёт net/margin
        try:
            obj.recompute()
        except Exception:
            c = float(getattr(obj, "cost", 0) or 0)
            p = float(getattr(obj, "price", 0) or 0)
            _safe_set(obj, "net", round(p - c, 2))
            _safe_set(obj, "margin", round(((p - c) / p) * 100, 1) if p else 0)

        # легаси-колонки, если есть
        _safe_set(obj, "sale_price_legacy", getattr(obj, "price", 0))
        _safe_set(obj, "net_profit_legacy", getattr(obj, "net", 0))

        db.session.add(obj)
        db.session.commit()
        flash("Создано", "success")
        return redirect(url_for("internal.tours_list"))

    if request.method == "POST" and form.errors:
        for fld, errs in form.errors.items():
            for e in errs:
                flash(f"{getattr(form, fld).label.text}: {e}", "danger")
        current_app.logger.error("internal.tour_create errors: %s", dict(form.errors))

    return render_template("internal/tour_form.html", form=form)


@bp.route("/tours/<int:tour_id>/edit", methods=["GET", "POST"], endpoint="edit_tour")
@login_required
def tour_edit(tour_id: int):
    _require_admin()
    obj = InternalTour.query.get_or_404(tour_id)
    form = InternalTourForm(obj=obj)

    # NEW: при GET подставляем коды в поля формы
    if request.method == "GET":
        form.supplier_code.data = (
            (obj.supplier.code if obj.supplier else getattr(obj, "supplier_code", None)) or ""
        )
        form.client_code.data = (
            (obj.client.code if obj.client else getattr(obj, "client_code", None)) or ""
        )

    if form.validate_on_submit():
        _safe_set(obj, "order_date", form.order_date.data, "date_order")
        _safe_set(obj, "start_date",  form.start_date.data, "date_start")
        _safe_set(obj, "end_date",    form.end_date.data,   "date_end")

        _safe_set(obj, "kind", form.kind.data)
        _safe_set(obj, "fio",  form.fio.data)
        _safe_set(obj, "direction", form.direction.data)

        # обновляем FK по кодам
        supplier_code = (form.supplier_code.data or "").strip()
        client_code   = (form.client_code.data or "").strip()
        supplier_id = _ensure_client_by_code(supplier_code)
        client_id   = _ensure_client_by_code(client_code)
        if not supplier_id or not client_id:
            flash("Укажите корректные коды поставщика и клиента.", "danger")
            return render_template("internal/tour_form.html", form=form, obj=obj)

        obj.supplier_id = supplier_id
        obj.client_id   = client_id

        _safe_set(obj, "cost",  form.cost.data or 0)
        _safe_set(obj, "price", form.price.data or 0)
        _safe_set(obj, "note",  form.note.data)

        try:
            obj.recompute()
        except Exception:
            c = float(getattr(obj, "cost", 0) or 0)
            p = float(getattr(obj, "price", 0) or 0)
            _safe_set(obj, "net", round(p - c, 2))
            _safe_set(obj, "margin", round(((p - c) / p) * 100, 1) if p else 0)

        _safe_set(obj, "sale_price_legacy", getattr(obj, "price", 0))
        _safe_set(obj, "net_profit_legacy", getattr(obj, "net", 0))

        db.session.commit()
        flash(f"Заказ #{obj.id} обновлён", "success")
        return redirect(url_for("internal.tours_list", **request.args.to_dict(flat=True)))

    if request.method == "POST" and form.errors:
        for fld, errs in form.errors.items():
            for e in errs:
                flash(f"{getattr(form, fld).label.text}: {e}", "danger")
        current_app.logger.error("internal.tour_edit errors: %s", dict(form.errors))

    return render_template("internal/tour_form.html", form=form, obj=obj)


@bp.route("/tours/<int:tour_id>/delete", methods=["POST", "GET"], endpoint="delete_tour")
@login_required
def tour_delete(tour_id: int):
    _require_admin()
    obj = InternalTour.query.get_or_404(tour_id)
    db.session.delete(obj)
    db.session.commit()
    flash(f"Заказ #{tour_id} удалён", "success")
    back = request.values.get("next") or url_for("internal.tours_list", **request.args.to_dict(flat=True))
    return redirect(back)
